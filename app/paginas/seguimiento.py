"""
PANTALLA DE SEGUIMIENTO
=======================
El tablero del supervisor y del master.

Los cálculos NO se hacen aquí: los hace la base, en las funciones
seguimiento_*. Eso tiene dos ventajas: la app no descarga miles de filas, y
el RLS aplica solo, así que el supervisor ve los mismos indicadores pero
limitados a sus zonas.
"""
from datetime import date, timedelta
from io import BytesIO

import pandas as pd
import streamlit as st

import config
import db

PERIODOS = {
    "Hoy": lambda: (date.today(), date.today()),
    "7 días": lambda: (date.today() - timedelta(days=6), date.today()),
    "Este mes": lambda: (date.today().replace(day=1), date.today()),
    "Mes pasado": lambda: _mes_pasado(),
}


def _mes_pasado():
    primero = date.today().replace(day=1)
    fin = primero - timedelta(days=1)
    return fin.replace(day=1), fin


def mostrar(usuario: dict):
    if usuario["rol"] not in ("supervisor", "master"):
        st.info("Esta pantalla es para supervisores y administradores.")
        return

    desde, hasta = _selector_periodo(usuario)

    resumen = db.seguimiento_resumen(desde, hasta)
    asesores = db.seguimiento_asesores(desde, hasta)
    zonas = db.seguimiento_zonas(desde, hasta)
    perdidas = db.seguimiento_perdidas(desde, hasta)
    ferreterias = db.seguimiento_ferreterias(desde, hasta)

    _pulso(resumen)
    _aperturas(desde, hasta)
    _tabla_asesores(asesores)

    col_zona, col_perdida = st.columns([1, 1.3])
    with col_zona:
        _tabla_zonas(zonas)
    with col_perdida:
        _tabla_perdidas(perdidas)

    _tabla_ferreterias(ferreterias)

    col_resumen, col_formato = st.columns(2)
    with col_resumen:
        _exportar(desde, hasta, resumen, asesores, zonas, perdidas, ferreterias)
    with col_formato:
        _exportar_formato_cotizadores(desde, hasta)


# ==================================================== selector de fechas ===
def _selector_periodo(usuario):
    col_titulo, col_rapido, col_fechas = st.columns([1.4, 1.6, 1.5],
                                                    vertical_alignment="bottom")

    with col_titulo:
        st.markdown("##### Seguimiento")
        zonas = [(z.get("m_zonas") or {}).get("nombre")
                 for z in (usuario.get("rel_usuarios_zonas") or [])]
        zonas = sorted(z for z in zonas if z)
        st.caption(", ".join(zonas) if zonas else "Todas las zonas")

    with col_rapido:
        atajo = st.segmented_control("Periodo", list(PERIODOS),
                                     default="Este mes",
                                     label_visibility="collapsed")

    if atajo:
        st.session_state.seguimiento_rango = PERIODOS[atajo]()

    rango_actual = st.session_state.get("seguimiento_rango") or PERIODOS["Este mes"]()

    with col_fechas:
        elegido = st.date_input("Rango", value=rango_actual,
                                format="DD/MM/YYYY", label_visibility="collapsed")

    # date_input devuelve una sola fecha mientras el usuario elige la segunda
    if isinstance(elegido, (list, tuple)) and len(elegido) == 2:
        st.session_state.seguimiento_rango = tuple(elegido)
        return elegido[0], elegido[1]

    return rango_actual[0], rango_actual[1]


# ============================================================ el pulso ===
def _pulso(r):
    if not r:
        st.info("No hay negociaciones en este periodo.")
        return

    c1, c2, c3, c4, c5, c6 = st.columns(6)
    c1.metric("Activas", r.get("activas", 0))
    c2.metric("Por cerrar", r.get("por_cerrar", 0))
    c3.metric("Ganadas", r.get("ganadas", 0))
    c4.metric("Perdidas", r.get("perdidas", 0))
    total = sum(r.get(k, 0) or 0 for k in ("activas", "por_cerrar", "ganadas", "perdidas"))
    c5.metric("Conversión", f"{r.get('conversion_pct') or 0}%",
              help=f"{r.get('ganadas', 0)} con venta de {total} negociaciones "
                   "abiertas en el periodo, sin importar si siguen activas.")
    c6.metric("Facturación", f"S/ {float(r.get('facturacion') or 0):,.0f}",
              help=f"{r.get('ventas', 0)} ventas registradas · "
                   f"{config.soles(r.get('facturacion'))}")

    por_validar = r.get("ventas_por_validar") or 0
    if por_validar:
        st.caption(f":orange[{por_validar} venta(s) esperando validación.] "
                   "Ya cuentan como ganadas, pero conviene confirmarlas.")


# ========================================================== evolutivo ===
def _aperturas(desde, hasta):
    filas = db.seguimiento_aperturas(desde, hasta)
    if not filas:
        return

    with st.container(border=True):
        total = sum(f["negociaciones"] for f in filas)
        promedio = round(total / len(filas), 1) if filas else 0
        st.markdown(f"**Apertura de negociaciones** · {total} en el periodo, "
                    f"promedio {promedio} por día")

        # El eje va como texto dd/mm: con fechas, Streamlit las rotula en
        # inglés y no hay forma de cambiarlo desde el gráfico.
        tabla = pd.DataFrame(filas)
        tabla["dia"] = pd.to_datetime(tabla["dia"]).dt.strftime("%d/%m")
        st.bar_chart(tabla.set_index("dia")["negociaciones"], height=200)


# =========================================================== por asesor ===
def _tabla_asesores(filas):
    with st.container(border=True):
        st.markdown("**Por asesor**")
        if not filas:
            st.caption("Sin datos en el periodo.")
            return

        tabla = pd.DataFrame([{
            "Asesor": f["asesor"],
            "Negociaciones": f["negociaciones"],
            "Cotizaciones": f["cotizaciones"],
            "Ventas": f["ventas"],
            "Facturación": float(f["facturacion"] or 0),
            "Conversión %": float(f["conversion_pct"] or 0),
            "Cotiz. por negoc.": float(f["cotiz_por_negociacion"] or 0),
        } for f in filas])

        st.dataframe(
            tabla, hide_index=True, width="stretch",
            column_config={
                "Facturación": st.column_config.NumberColumn(format="S/ %.2f"),
                "Conversión %": st.column_config.ProgressColumn(
                    format="%.1f%%", min_value=0, max_value=100),
                "Cotiz. por negoc.": st.column_config.NumberColumn(
                    format="%.1f",
                    help="Recotizar mucho y cerrar poco señala un problema de "
                         "precio o de calificación del cliente."),
            },
        )


# ============================================================= por zona ===
def _tabla_zonas(filas):
    with st.container(border=True):
        st.markdown("**Por zona**")
        if not filas:
            st.caption("Sin datos en el periodo.")
            return

        st.dataframe(pd.DataFrame([{
            "Zona": f["zona"],
            "Negociaciones": f["negociaciones"],
            "Ventas": f["ventas"],
            "Facturación": float(f["facturacion"] or 0),
            "Conversión %": float(f["conversion_pct"] or 0),
        } for f in filas]), hide_index=True, width="stretch",
            column_config={
                "Facturación": st.column_config.NumberColumn(format="S/ %.2f"),
                "Conversión %": st.column_config.NumberColumn(format="%.1f%%"),
            })


# ======================================================== por qué pierden ===
def _tabla_perdidas(filas):
    with st.container(border=True):
        st.markdown("**Por qué se pierden**")
        if not filas:
            st.caption("Ninguna negociación perdida en el periodo.")
            return

        st.dataframe(pd.DataFrame([{
            "Motivo": f["motivo"] + (" (automático)" if f["automatico"] else ""),
            "Casos": f["casos"],
            "Monto": float(f["monto"] or 0),
        } for f in filas]), hide_index=True, width="stretch",
            column_config={"Monto": st.column_config.NumberColumn(format="S/ %.2f")})

        st.caption("El monto sale de la última cotización de cada negociación.")


# ========================================================== ferreterías ===
def _tabla_ferreterias(filas):
    with st.container(border=True):
        st.markdown("**Ferreterías**")
        if not filas:
            st.caption("Sin cotizaciones en el periodo.")
            return

        tabla = pd.DataFrame([{
            "Ferretería": f["ferreteria"],
            "Compitió": f["compitio"],
            "Elegida": f["elegida"],
            "Vendida": f["vendida"],
            "% elegida": float(f["tasa_elegida_pct"] or 0),
            "% cierre": float(f["conversion_pct"] or 0),
        } for f in filas])

        st.dataframe(tabla, hide_index=True, width="stretch",
                     column_config={
                         "% elegida": st.column_config.ProgressColumn(
                             format="%.1f%%", min_value=0, max_value=100,
                             help="De las veces que compitió, cuántas la eligió el asesor."),
                         "% cierre": st.column_config.NumberColumn(format="%.1f%%"),
                     })

        # La señal más accionable del tablero
        flojas = [f for f in filas
                  if f["compitio"] >= 10 and (f["tasa_elegida_pct"] or 0) < 15]
        if flojas:
            nombres = ", ".join(f["ferreteria"] for f in flojas[:3])
            st.caption(f":orange[{nombres} compite seguido y casi nunca es "
                       "elegida: sus precios no son competitivos.]")


# ============================================================== exportar ===
def _exportar(desde, hasta, resumen, asesores, zonas, perdidas, ferreterias):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame([resumen]).to_excel(writer, sheet_name="Resumen", index=False)
        for nombre, filas in [("Asesores", asesores), ("Zonas", zonas),
                              ("Perdidas", perdidas), ("Ferreterias", ferreterias)]:
            pd.DataFrame(filas).to_excel(writer, sheet_name=nombre, index=False)

    st.download_button(
        "⬇ Exportar a Excel",
        data=buffer.getvalue(),
        file_name=f"seguimiento_{desde}_{hasta}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
    )


# ================================================= formato histórico ===
def _exportar_formato_cotizadores(desde, hasta):
    """Cotizaciones y ventas en el formato del cotizador anterior.

    Una fila por producto. ID 3-n es cotización y 4-n es venta; en las
    ventas, COTIZACION_ORIGEN dice de qué cotización salieron.
    """
    # Se genera solo cuando se pide: con muchos meses pueden ser miles de
    # filas y no vale la pena consultarlas en cada recarga de la pantalla.
    clave = f"formato_{desde}_{hasta}"
    if st.session_state.get("formato_clave") != clave:
        st.session_state.formato_archivo = None

    if st.session_state.get("formato_archivo") is None:
        if st.button("Preparar formato cotizadores", width="stretch",
                     help="Cotizaciones y ventas del periodo, una fila por producto, "
                          "con las columnas del cotizador anterior."):
            with st.spinner("Armando el archivo…"):
                filas = db.formato_cotizadores(desde, hasta)
                st.session_state.formato_archivo = _excel_formato(filas)
                st.session_state.formato_filas = len(filas)
                st.session_state.formato_clave = clave
            st.rerun()
        return

    st.download_button(
        f"⬇ Formato cotizadores · {st.session_state.formato_filas:,} filas",
        data=st.session_state.formato_archivo,
        file_name=f"formato_cotizadores_{desde}_{hasta}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
    )


# El orden exacto del formato histórico. Se fija aquí para no depender del
# orden en que lleguen los datos.
COLUMNAS_FORMATO = [
    "ID", "FECHA", "CLIENTE", "DOCUMENTO", "DEPARTAMENTO", "PROVINCIA",
    "DISTRITO", "ITEM", "Categoria", "Marca", "Producto", "id", "Um",
    "Precio", "Cantidad", "Precio Total", "TC", "DISTRIBUIDOR",
    "PRECIO_FECO", "TIPO_CLIENTE", "ACCION", "TELEFONO", "TIPO_PAGO",
    "TICKET", "COTIZACION_ORIGEN",
]


def _excel_formato(filas) -> bytes:
    tabla = pd.DataFrame(filas)
    if tabla.empty:
        tabla = pd.DataFrame(columns=COLUMNAS_FORMATO)
    else:
        # El ID es texto: ordenado como texto quedaría 3-1, 3-10, 3-2.
        # Por eso se ordena con la columna auxiliar, que después no se exporta.
        tabla = tabla.sort_values(["FECHA", "ACCION", "orden_documento", "ITEM"])
        tabla = tabla[COLUMNAS_FORMATO]
        tabla["FECHA"] = pd.to_datetime(tabla["FECHA"]).dt.date

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        tabla.to_excel(writer, sheet_name="Formato", index=False)
        hoja = writer.sheets["Formato"]
        hoja.freeze_panes = "A2"
        for celda in hoja[1]:
            celda.font = celda.font.copy(bold=True)
        for i, columna in enumerate(tabla.columns, start=1):
            ancho = max(10, min(40, len(str(columna)) + 4))
            if columna in ("CLIENTE", "Producto", "id", "DISTRITO"):
                ancho = 28
            from openpyxl.utils import get_column_letter
            hoja.column_dimensions[get_column_letter(i)].width = ancho
        # Fecha en formato peruano
        for fila in hoja.iter_rows(min_row=2, min_col=2, max_col=2):
            for celda in fila:
                celda.number_format = "DD/MM/YYYY"
    return buffer.getvalue()
